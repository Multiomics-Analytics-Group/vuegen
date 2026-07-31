"""Tests for ensuring Quarto cell labels are unique when files share the same stem."""

import re
from pathlib import Path

import openpyxl

from vuegen.report_generator import get_report


def _get_labels_from_qmd(qmd_path: Path) -> list:
    """Extract all cell labels from a QMD file."""
    content = qmd_path.read_text()
    return re.findall(r"#\| label: '(.+?)'", content)


def test_unique_labels_with_same_stem_files(tmp_path):
    """
    Test that Quarto cell labels are unique when files with the same stem but
    different extensions (csv, xlsx, md) exist in the same directory.

    Regression test for: https://github.com/Multiomics-Analytics-Group/vuegen/issues/192
    """
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    output_dir = tmp_path / "output"
    output_dir.mkdir()

    # Create files with the same name stem 'metadata_samples'
    (data_dir / "metadata_samples.csv").write_text("col1,col2\n1,2\n3,4")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "col1"
    ws["B1"] = "col2"
    wb.save(data_dir / "metadata_samples.xlsx")

    (data_dir / "metadata_samples.md").write_text("# Metadata Samples\nSome content.")

    report_dir, _ = get_report(
        "html",
        dir_path=str(data_dir),
        output_dir=output_dir,
    )

    qmd_path = output_dir / "quarto_report" / "quarto_report.qmd"
    assert qmd_path.exists(), f"QMD file not found at {qmd_path}"

    labels = _get_labels_from_qmd(qmd_path)
    assert len(labels) == len(set(labels)), (
        f"Duplicate Quarto cell labels found: "
        f"{[l for l in labels if labels.count(l) > 1]}"
    )


def test_unique_labels_no_collision(tmp_path):
    """
    Test that labels are not modified when there are no collisions.

    Existing behavior should be preserved for the normal case.
    """
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    output_dir = tmp_path / "output"
    output_dir.mkdir()

    # Create files with different stems (no collision)
    (data_dir / "samples.csv").write_text("col1,col2\n1,2\n3,4")
    (data_dir / "readme.md").write_text("# Readme\nSome content.")

    report_dir, _ = get_report(
        "html",
        dir_path=str(data_dir),
        output_dir=output_dir,
    )

    qmd_path = output_dir / "quarto_report" / "quarto_report.qmd"
    assert qmd_path.exists(), f"QMD file not found at {qmd_path}"

    labels = _get_labels_from_qmd(qmd_path)
    assert len(labels) == len(set(labels)), (
        f"Duplicate Quarto cell labels found: "
        f"{[l for l in labels if labels.count(l) > 1]}"
    )

    # Labels should not have collision-resolution suffixes for non-duplicate cases
    # A collision suffix would be like 'Samples 1-2' (numeric-numeric pattern)
    multi_numeric_labels = [l for l in labels if re.search(r"\d+-\d+$", l)]
    assert len(multi_numeric_labels) == 0, (
        f"Unexpected collision-resolution suffixes in labels: {multi_numeric_labels}"
    )
